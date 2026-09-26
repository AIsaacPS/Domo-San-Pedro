"""
extraer_texto_bibliografia.py
=============================
Extrae el texto de TODA la bibliografía (PDF y DOCX) de la carpeta BIBLIOGRAFIA/
a archivos .txt en docs/extracted-text/, para que puedan ser leídos por los
agentes auditores y usados como sustento en la redacción de la tesina.

DISEÑO: descubre los archivos AUTOMÁTICAMENTE recorriendo BIBLIOGRAFIA/ de forma
recursiva. No hay lista fija de rutas, así que reorganizar las carpetas no rompe
el script. Los .txt de salida conservan la estructura de subcarpetas por tema.

USO:
    python scripts/extraer_texto_bibliografia.py

REQUISITOS: pymupdf (fitz) para PDF; python-docx para DOCX.

NOTA: Se ignora la subcarpeta "MAPAS IGNORAR" y los .txt ya existentes en
      BIBLIOGRAFIA (p. ej. el propio capítulo de la tesina).
"""

from pathlib import Path
import fitz  # pymupdf

try:
    import docx  # python-docx
    HAY_DOCX = True
except ImportError:
    HAY_DOCX = False

PROJECT_DIR = Path(__file__).parent.parent
BIBLIO_DIR = PROJECT_DIR / "BIBLIOGRAFIA"
OUT_DIR = PROJECT_DIR / "docs" / "extracted-text"
OUT_DIR.mkdir(parents=True, exist_ok=True)

# Carpetas cuyo contenido se ignora (por nombre, en cualquier nivel).
IGNORAR_DIRS = {"MAPAS IGNORAR"}


def slugify(texto):
    """Nombre de archivo .txt limpio (sin espacios ni acentos)."""
    reemplazos = {" ": "_", "–": "-", "á": "a", "é": "e", "í": "i",
                  "ó": "o", "ú": "u", "ñ": "n", "Á": "A", "É": "E",
                  "Í": "I", "Ó": "O", "Ú": "U", "Ñ": "N", ",": ""}
    for a, b in reemplazos.items():
        texto = texto.replace(a, b)
    return "".join(c for c in texto if c.isalnum() or c in "._-")


def salida_para(archivo):
    """Ruta .txt de salida que conserva la subcarpeta temática de origen."""
    rel = archivo.relative_to(BIBLIO_DIR)
    sub = "_".join(slugify(p) for p in rel.parts[:-1])  # carpetas temáticas
    nombre = slugify(rel.stem) + ".txt"
    if sub:
        nombre = f"{sub}__{nombre}"
    return OUT_DIR / nombre


def extraer_pdf(path):
    doc = fitz.open(str(path))
    partes = [page.get_text() for page in doc]
    doc.close()
    return "\n".join(partes)


def extraer_docx(path):
    d = docx.Document(str(path))
    partes = [p.text for p in d.paragraphs]
    # también el texto de las tablas
    for tabla in d.tables:
        for fila in tabla.rows:
            partes.append("\t".join(celda.text for celda in fila.cells))
    return "\n".join(partes)


def descubrir_archivos():
    """Lista todos los PDF/DOCX de BIBLIOGRAFIA, saltando carpetas ignoradas."""
    archivos = []
    for p in sorted(BIBLIO_DIR.rglob("*")):
        if p.is_dir():
            continue
        if any(parte in IGNORAR_DIRS for parte in p.parts):
            continue
        if p.suffix.lower() in (".pdf", ".docx"):
            archivos.append(p)
    return archivos


def main():
    print("=" * 66)
    print("  EXTRACCIÓN DE TEXTO — Bibliografía Domo San Pedro")
    print("=" * 66)

    archivos = descubrir_archivos()
    print(f"  Encontrados: {len(archivos)} archivos (PDF/DOCX)\n")

    resumen = []
    for f in archivos:
        rel = f.relative_to(BIBLIO_DIR)
        try:
            if f.suffix.lower() == ".pdf":
                texto = extraer_pdf(f)
            elif f.suffix.lower() == ".docx":
                if not HAY_DOCX:
                    print(f"  ! [SIN python-docx] {rel}")
                    resumen.append((str(rel), "SIN_DOCX", 0))
                    continue
                texto = extraer_docx(f)
            else:
                continue
        except Exception as e:
            print(f"  ! [ERROR] {rel}: {e}")
            resumen.append((str(rel), "ERROR", 0))
            continue

        out = salida_para(f)
        out.write_text(texto, encoding="utf-8")
        n = len(texto.strip())
        estado = "OK" if n > 200 else "VACIO/ESCANEADO"
        print(f"  [{estado}] {out.name}  ({n:,} chars)")
        resumen.append((str(rel), estado, n))

    print("\n" + "=" * 66)
    ok = sum(1 for _, e, _ in resumen if e == "OK")
    problemas = [(r, e) for r, e, _ in resumen if e != "OK"]
    print(f"  RESUMEN: {ok}/{len(resumen)} extraídos con texto limpio")
    if problemas:
        print("  Requieren atención:")
        for r, e in problemas:
            print(f"    - [{e}] {r}")
    print(f"  Salida: {OUT_DIR}")
    print("=" * 66)


if __name__ == "__main__":
    main()
